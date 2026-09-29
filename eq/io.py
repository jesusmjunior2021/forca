"""Leitura do arquivo ORIGINAL (xlsx/csv) como matriz crua e exportações (xlsx, csv, pdf, matriz formatada)."""
import io, datetime
import pandas as pd
from openpyxl import load_workbook, Workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

NAVY = "1F3864"


def ler_matriz(conteudo: bytes, nome: str):
    if nome.lower().endswith(".csv"):
        for enc in ("utf-8-sig", "latin-1"):
            try:
                df = pd.read_csv(io.BytesIO(conteudo), header=None, dtype=object, sep=None, engine="python", encoding=enc)
                break
            except UnicodeDecodeError:
                continue
        return [[None if pd.isna(c) else c for c in r] for r in df.values.tolist()]
    wb = load_workbook(io.BytesIO(conteudo), data_only=True, read_only=False)
    ws = wb.worksheets[0]
    rows = [[c for c in r] for r in ws.iter_rows(values_only=True)]
    while rows and all(c is None for c in rows[-1]):
        rows.pop()
    return rows


ROTULOS = {"universo": "Universo", "rank": "Rank IPCO", "lot": "Lotação", "unidade": "Unidade / Vara", "comarca": "Comarca", "entrancia": "Entrância",
           "tipo": "Tipo de unidade", "prev": "Prevista", "pre": "Preenchido", "deficit": "Déficit (+falta/−sobra)", "situacao": "Situação",
           "media": "Média casos novos", "tx": "Tx congestionamento", "afast": "Afastados >60d", "Nd": "N_Déficit", "Nm": "N_Média", "Nt": "N_TxCong",
           "ipco": "IPCO [0-1]", "classe": "Classe", "semDados": "Sem dados estatísticos", "recente": "Instalada recentemente",
           "cargaServidor": "Carga (casos novos / servidor)", "res": "Residentes", "est": "Estagiários", "solic": "Solicitação Resid./Estag."}


def df_rotulado(df):
    return df[[c for c in ROTULOS if c in df.columns]].rename(columns=ROTULOS)


def xlsx_bytes(df, nome_aba="Dados", largura=None):
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="xlsxwriter") as w:
        df.to_excel(w, sheet_name=nome_aba[:31], index=False)
        ws, wb = w.sheets[nome_aba[:31]], w.book
        h = wb.add_format({"bold": True, "font_color": "white", "bg_color": "#" + NAVY, "text_wrap": True, "valign": "vcenter", "font_name": "Arial"})
        for i, c in enumerate(df.columns):
            ws.write(0, i, c, h)
            ws.set_column(i, i, min(max(len(str(c)) + 2, int(df[c].astype(str).str.len().quantile(0.9)) + 2 if len(df) else 10), 55))
        ws.freeze_panes(1, 0); ws.autofilter(0, 0, len(df), len(df.columns) - 1)
    return buf.getvalue()


def pdf_bytes(titulo, subtitulo, df):
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet
    buf = io.BytesIO(); st_ = getSampleStyleSheet()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=24, rightMargin=24, topMargin=24, bottomMargin=24)
    cols = ["Rank IPCO", "Lotação", "Unidade / Vara", "Comarca", "Entrância", "Prevista", "Preenchido", "Déficit (+falta/−sobra)", "IPCO [0-1]", "Classe", "Residentes", "Estagiários"]
    cols = [c for c in cols if c in df.columns]
    fmt = lambda v: "" if pd.isna(v) else (f"{v:.3f}" if isinstance(v, float) and not float(v).is_integer() else (str(int(v)) if isinstance(v, float) else str(v)))
    small = st_["BodyText"].clone("s", fontSize=6.5, leading=8)
    data = [[Paragraph(f"<b>{c}</b>", small) for c in cols]] + [[Paragraph(fmt(r[c])[:70], small) for c in cols] for _, r in df.iterrows()]
    t = Table(data, repeatRows=1, colWidths=[None] * len(cols))
    t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#" + NAVY)), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                           ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#BFBFBF")), ("VALIGN", (0, 0), (-1, -1), "MIDDLE")]))
    for i, c in enumerate(df["Classe"].tolist() if "Classe" in df else [], start=1):
        cor = {"ALTA": "#F4B6B6", "MÉDIA": "#FCD9A8", "BAIXA": "#C6E7C6"}.get(c)
        if cor: t.setStyle(TableStyle([("BACKGROUND", (cols.index("Classe"), i), (cols.index("Classe"), i), colors.HexColor(cor))]))
    doc.build([Paragraph(titulo, st_["Title"]), Paragraph(subtitulo, st_["Normal"]), Spacer(1, 8), t])
    return buf.getvalue()


# ---------- Matriz de CRITICIDADE formatada (mesmo layout do arquivo final) ----------
_LEG = ("● IPCO ≥ 0,70 → Criticidade ALTA (vermelho)    ● IPCO 0,40–0,69 → Criticidade MÉDIA (laranja)    ● IPCO < 0,40 → Criticidade BAIXA (verde)    "
        "● Texto em negrito → unidade com Residente(s) e/ou Estagiário(s)    ● N_Déficit com sinal: (+) falta de servidores | (−) excedente ao quadro    "
        "● Escala bipartida: superávit reduz a contribuição do critério Déficit abaixo do equilíbrio    ● Fundo cinza → sem dados estatísticos")
_META = ("Metodologia: escala bipartida centrada em zero para o Déficit (superávit −1 a 0 | equilíbrio 0 | déficit 0 a +1), pesos Déficit 60% + Média de Casos Novos 30% + "
         "Tx. Congestionamento 10%. IPCO reescalonado para [0,1] pelo máximo e mínimo do conjunto, garantindo que a unidade de maior criticidade resulte em 1,0000 e a de menor em 0,0000. "
         "N_Déficit com sinal: positivo = falta | negativo = excedente.")
_META_SEJ = ("Metodologia: escala bipartida centrada em zero para o Déficit (superávit −1 a 0 | equilíbrio 0 | déficit 0 a +1), pesos 50% Déficit + 50% Média de Casos Novos. "
             "IPCO reescalonado para [0,1] pelo máximo e mínimo do conjunto. N_Déficit com sinal: positivo = falta | negativo = excedente ao quadro.")
_COLS = ["Rank\nIPCO", "Lotação", "Unidade / Vara", "Comarca", "Entrância", "Prevista (efetivo)", "Preenchido", "Déficit\n(+falta/−sobra)", "Média Casos\nNovos", "Tx.\nCong.",
         "N_Déficit\n(norm.)", "N_Média\n(norm.)", "N_Txcong\n(norm.)", "IPCO\n[0-1]", "Residentes", "Estagiários", "Solicitação Resid./Estag."]
_FMT = ["0", "0", "@", "@", "@", "0", "0", "0", "#,##0", "0.00%", "+0.0000;-0.0000;0.0000", "0.0000", "0.0000", "0.0000", "0", "0", "@"]
_DEF = [("SECRE", "IPCO - Rank Criticidade - SECRE", "ÍNDICE PRELIMINAR DE CRITICIDADE OPERACIONAL (IPCO) — CARGOS DE SECRETARIA (secretaria)", _META, False),
        ("GAB", "IPCO - Rank Criticidade - GAB", "ÍNDICE PRELIMINAR DE CRITICIDADE OPERACIONAL (IPCO) — CARGOS DE GABINETE (estagiário)", _META, True),
        ("SEJUD", "Unidades Atendidas SEJUD", "UNIDADES JUDICIAIS ATENDIDAS PELAS SECRETARIAS JUDICIAIS ÚNICAS DIGITAIS (SEJUD)", _META_SEJ, True)]
_COR = {"ALTA": "FF4444", "MÉDIA": "FFA500", "BAIXA": "5CB85C"}


def matriz_formatada(df) -> bytes:
    wb = Workbook(); wb.remove(wb.active)
    thin = Side(style="thin", color="BFBFBF"); box = Border(left=thin, right=thin, top=thin, bottom=thin)
    for u, nome, titulo, meta, solic in _DEF:
        d = df[df["universo"] == u].sort_values("rank")
        if d.empty:
            continue
        nc = 17 if solic else 16
        ws = wb.create_sheet(nome)
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=nc); ws.cell(1, 1, titulo)
        ws["A1"].font = Font(name="Arial", bold=True, size=13, color="FFFFFF"); ws["A1"].fill = PatternFill("solid", fgColor=NAVY); ws.row_dimensions[1].height = 30
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=nc); ws.cell(2, 1, meta)
        ws["A2"].font = Font(name="Arial", italic=True, size=9); ws["A2"].alignment = Alignment(wrap_text=True, vertical="top"); ws.row_dimensions[2].height = 58
        for i, h in enumerate(_COLS[:nc], start=1):
            c = ws.cell(4, i, h); c.font = Font(name="Arial", bold=True, color="FFFFFF", size=10); c.fill = PatternFill("solid", fgColor=NAVY)
            c.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center"); c.border = box
        ws.row_dimensions[4].height = 42
        n = lambda v: "—" if v is None or (isinstance(v, float) and v != v) else v
        for k, r in enumerate(d.to_dict("records")):
            row = 5 + k
            vals = [r["rank"], r["lot"], r["unidade"], r["comarca"], r["entrancia"] or "—", n(r["prev"]), n(r["pre"]), n(r["deficit"]), n(r["media"]), n(r["tx"]),
                    n(r["Nd"]), n(r["Nm"]), n(r["Nt"]), n(r["ipco"]), n(r["res"]), n(r["est"])] + ([r["solic"] or ""] if solic else [])
            neg = (r["res"] or 0) > 0 or (r["est"] or 0) > 0
            for i, v in enumerate(vals, start=1):
                c = ws.cell(row, i, v); c.number_format = _FMT[i - 1]; c.font = Font(name="Arial", size=10, bold=neg); c.border = box
                if r["semDados"] == "SIM": c.fill = PatternFill("solid", fgColor="EDEDED")
            if r["classe"] in _COR:
                for i in (1, 14):
                    c = ws.cell(row, i); c.fill = PatternFill("solid", fgColor=_COR[r["classe"]]); c.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        fim = 4 + len(d)
        ws.auto_filter.ref = f"A4:{get_column_letter(nc)}{fim}"; ws.freeze_panes = "D5"
        ws.merge_cells(start_row=fim + 2, start_column=1, end_row=fim + 2, end_column=nc); ws.cell(fim + 2, 1, _LEG)
        ws.cell(fim + 2, 1).font = Font(name="Arial", size=9); ws.cell(fim + 2, 1).alignment = Alignment(wrap_text=True, vertical="top"); ws.row_dimensions[fim + 2].height = 48
        for i, w in enumerate([8, 10, 55, 28, 15, 11, 11, 12, 12, 9, 11, 11, 11, 9, 11, 11, 60][:nc], start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()
